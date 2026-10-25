"""Build the revised manuscript: ms_src.md -> ms_final.md -> .docx (marked and clean)."""
import re, json, subprocess, sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

FIGDIR = "/home/claude/nb/full/results"
NEWREF = {"rosmanis23", "rosmanis24", "stoudenmire", "leng", "sahar", "menu", "dowarah", "abughanem", "qiskit"}
REFS = {
 "feynman": "Feynman, R.P. Simulating physics with computers. *Int. J. Theor. Phys.* **1982**, *21*, 467–488. https://doi.org/10.1007/BF02650179.",
 "aaronson": "Aaronson, S. *Quantum Computing since Democritus*; Cambridge University Press: Cambridge, UK, 2013.",
 "dj": "Deutsch, D.; Jozsa, R. Rapid solution of problems by quantum computation. *Proc. R. Soc. Lond. A* **1992**, *439*, 553–558. https://doi.org/10.1098/rspa.1992.0167.",
 "shor": "Shor, P.W. Polynomial-time algorithms for prime factorization and discrete logarithms on a quantum computer. *SIAM J. Comput.* **1997**, *26*, 1484–1509. https://doi.org/10.1137/S0097539795293172.",
 "grover": "Grover, L.K. A fast quantum mechanical algorithm for database search. In Proceedings of the 28th Annual ACM Symposium on Theory of Computing, Philadelphia, PA, USA, 22–24 May 1996; pp. 212–219. https://doi.org/10.1145/237814.237866.",
 "bhmt": "Brassard, G.; Høyer, P.; Mosca, M.; Tapp, A. Quantum amplitude amplification and estimation. *Contemp. Math.* **2002**, *305*, 53–74. https://doi.org/10.1090/conm/305/05215.",
 "cleve": "Cleve, R.; Ekert, A.; Macchiavello, C.; Mosca, M. Quantum algorithms revisited. *Proc. R. Soc. Lond. A* **1998**, *454*, 339–354. https://doi.org/10.1098/rspa.1998.0164.",
 "bbht": "Boyer, M.; Brassard, G.; Høyer, P.; Tapp, A. Tight bounds on quantum searching. *Fortschr. Phys.* **1998**, *46*, 493–505.{! [VERIFY: add DOI]!}",
 "zalka": "Zalka, C. Grover’s quantum searching algorithm is optimal. *Phys. Rev. A* **1999**, *60*, 2746–2751. https://doi.org/10.1103/PhysRevA.60.2746.",
 "beals": "Beals, R.; Buhrman, H.; Cleve, R.; Mosca, M.; de Wolf, R. Quantum lower bounds by polynomials. *J. ACM* **2001**, *48*, 778–797. https://doi.org/10.1145/502090.502097.",
 "bbbv": "Bennett, C.H.; Bernstein, E.; Brassard, G.; Vazirani, U. Strengths and weaknesses of quantum computing. *SIAM J. Comput.* **1997**, *26*, 1510–1523. https://doi.org/10.1137/S0097539796300933.",
 "preskill18": "Preskill, J. Quantum computing in the NISQ era and beyond. *Quantum* **2018**, *2*, 79. https://doi.org/10.22331/q-2018-08-06-79.",
 "montanaro": "Montanaro, A. Quantum algorithms: An overview. *npj Quantum Inf.* **2016**, *2*, 15023. https://doi.org/10.1038/npjqi.2015.23.",
 "arute": "Arute, F.; Arya, K.; Babbush, R.; Bacon, D.; Bardin, J.C.; Barends, R.; et al. Quantum supremacy using a programmable superconducting processor. *Nature* **2019**, *574*, 505–510. https://doi.org/10.1038/s41586-019-1666-5.",
 "preskill98": "Preskill, J. Fault-tolerant quantum computation. In *Introduction to Quantum Computation and Information*; Lo, H.-K., Popescu, S., Spiller, T., Eds.; World Scientific: Singapore, 1998; pp. 213–269.",
 "kitaev": "Kitaev, A.Y. Fault-tolerant quantum computation by anyons. *Ann. Phys.* **2003**, *303*, 2–30. https://doi.org/10.1016/S0003-4916(02)00018-0.",
 "pablo": "Pablo-Norman, B.; Ruiz-Altaba, M. Noise in Grover’s quantum search algorithm. *Phys. Rev. A* **1999**, *61*, 012301. https://doi.org/10.1103/PhysRevA.61.012301.",
 "long00": "Long, G.L.; Li, Y.S.; Zhang, W.L.; Tu, C.C. Dominant gate imperfection in Grover’s quantum search algorithm. *Phys. Rev. A* **2000**, *61*, 042305. https://doi.org/10.1103/PhysRevA.61.042305.",
 "shenvi": "Shenvi, N.; Brown, K.R.; Whaley, K.B. Effects of a random noisy oracle on search algorithm complexity. *Phys. Rev. A* **2003**, *68*, 052313. https://doi.org/10.1103/PhysRevA.68.052313.",
 "regev": "Regev, O.; Schiff, L. Impossibility of a quantum speed-up with a faulty oracle. In *Automata, Languages and Programming (ICALP 2008)*; Lecture Notes in Computer Science, Vol. 5125; Springer: Berlin/Heidelberg, Germany, 2008; pp. 773–781.{! [VERIFY: add DOI]!}",
 "rosmanis23": "Rosmanis, A. Quantum search with noisy oracle. *arXiv* **2023**, arXiv:2309.14944.",
 "rosmanis24": "Rosmanis, A. Addendum to “Quantum search with noisy oracle”. *arXiv* **2024**, arXiv:2405.11973.",
 "stoudenmire": "Stoudenmire, E.M.; Waintal, X. Opening the black box inside Grover’s algorithm. *Phys. Rev. X* **2024**{! [VERIFY: volume, article number, DOI]!}; arXiv:2303.11317.",
 "leng": "Leng, J.; Yang, F.; Wang, X.-B. Noise-tolerant Grover’s algorithm via success-probability prediction. *Phys. Rev. Res.* **2025**, *7*, L012017. https://doi.org/10.1103/PhysRevResearch.7.L012017.",
 "sahar": "Saha Roy, A.; Lacalle, J. Effect of isotropic errors on the complexity of Grover’s algorithm. *arXiv* **2026**, arXiv:2606.04817.",
 "menu": "Menu, R.; Schachenmayer, J. Efficiency of classical simulations of a noisy Grover algorithm. *arXiv* **2026**, arXiv:2602.17569.",
 "dowarah": "Dowarah, S.; Zhang, C.; Khemani, V.; Kolodrubetz, M. Phases and phase transition in Grover’s algorithm with systematic noise. *arXiv* **2024**, arXiv:2406.10344.{! [VERIFY: journal version]!}",
 "abughanem": "AbuGhanem, M. Characterizing Grover search algorithm on large-scale superconducting quantum computers. *Sci. Rep.* **2025**, *15*, 1281. https://doi.org/10.1038/s41598-024-80188-6.",
 "long01": "Long, G.L. Grover algorithm with zero theoretical failure rate. *Phys. Rev. A* **2001**, *64*, 022307. https://doi.org/10.1103/PhysRevA.64.022307.",
 "toirov24": "Toirov, Sh.A.; Mingboyev, U.X.; Parmonov, A.A.; Kudratov, R.B. Designing Grover’s quantum searching algorithm for 2-qubit and 3-qubit states of quantum systems. In Proceedings of the III International Scientific and Practical Conference “Technologies, Materials Science and Engineering”, *AIP Conf. Proc.* **2024**, *3243*, 010001.{! [VERIFY: article number/DOI]!}",
 "toirov21": "Toirov, Sh.; Narmuradov, U. Processes of applying of quantum genetic algorithm in function optimization. *AIP Conf. Proc.* **2021**, *2365*, 1.{! [VERIFY: article number/DOI]!}",
 "nc": "Nielsen, M.A.; Chuang, I.L. *Quantum Computation and Quantum Information*, 10th Anniversary ed.; Cambridge University Press: Cambridge, UK, 2010. https://doi.org/10.1017/CBO9780511976667.",
 "barenco": "Barenco, A.; Bennett, C.H.; Cleve, R.; DiVincenzo, D.P.; Margolus, N.; Shor, P.; Sleator, T.; Smolin, J.A.; Weinfurter, H. Elementary gates for quantum computation. *Phys. Rev. A* **1995**, *52*, 3457–3467. https://doi.org/10.1103/PhysRevA.52.3457.",
 "qiskit": "Javadi-Abhari, A.; et al. Quantum computing with Qiskit. *arXiv* **2024**, arXiv:2405.08810.{! [VERIFY: full author list per MDPI style]!}",
 "qaoa": "Farhi, E.; Goldstone, J.; Gutmann, S. A quantum approximate optimization algorithm. *arXiv* **2014**, arXiv:1411.4028.",
}

def esc_outside_math(t):
    parts = re.split(r"(\$\$.*?\$\$|\$[^$]*\$|`[^`]*`)", t, flags=re.S)
    return "".join(p if (p.startswith("$") or p.startswith("`")) else p.replace("[", "\\[").replace("]", "\\]") for p in parts)

def build(marked=True):
    s = open("ms_src.md").read()
    T = json.load(open("tables.json"))
    for k in ["T7", "T7b", "T8", "T9", "T10", "T11"]:
        s = s.replace("{{" + k + "}}", T[k])
    s = s.replace("FIGDIR", FIGDIR)
    # citations in order of first appearance
    order = []
    def cite(m):
        keys = [k.strip() for k in m.group(1).split(",")]
        for k in keys:
            assert k in REFS, k
            if k not in order: order.append(k)
        nums = sorted(order.index(k) + 1 for k in keys)
        out, i = [], 0
        while i < len(nums):
            j = i
            while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1: j += 1
            out.append(f"{nums[i]}–{nums[j]}" if j - i >= 2 else ",".join(map(str, nums[i:j + 1])))
            i = j + 1
        return "⟦" + ",".join(out) + "⟧"
    s = re.sub(r"\{c:([^}]+)\}", cite, s)
    # NEW> paragraphs
    s = re.sub(r"^NEW> (.*)$", lambda m: "{+" + m.group(1) + "+}", s, flags=re.M)
    # references list
    reflines = []
    for i, k in enumerate(order, 1):
        r = REFS[k]
        if k in NEWREF:
            main, _, ver = r.partition("{!")
            r = "{+" + main + "+}" + ("{!" + ver if ver else "")
        reflines.append(f"{i}. " + r)
    s = s.replace("REFLIST", "\n".join(reflines))
    # spans
    s = re.sub(r"\{!(.*?)!\}", lambda m: "[" + esc_outside_math(m.group(1)) + ']{custom-style="Verify Text"}', s, flags=re.S)
    s = re.sub(r"\{\+(.*?)\+\}", lambda m: "[" + esc_outside_math(m.group(1)) + ']{custom-style="New Text"}', s, flags=re.S)
    s = s.replace("⟦", "\\[").replace("⟧", "\\]")
    open("ms_final.md", "w").write(s)
    unused = [k for k in REFS if k not in order]
    print("references:", len(order), "unused keys:", unused)
    return order

order = build()
subprocess.run(["pandoc", "ms_final.md", "-f", "markdown+tex_math_dollars+bracketed_spans+fenced_divs+pipe_tables+implicit_figures",
                "-o", "_raw.docx"], check=True)

def style_doc(src, dst, marked):
    d = Document(src)
    st = d.styles
    def font(name, size=None, bold=None, color=None, italic=None, align=None, space_after=None):
        try:
            x = st[name]
        except KeyError:
            return
        f = x.font
        f.name = "Palatino Linotype"
        if size: f.size = Pt(size)
        if bold is not None: f.bold = bold
        if italic is not None: f.italic = italic
        if color is not None: f.color.rgb = color
        if align is not None and hasattr(x, "paragraph_format"): x.paragraph_format.alignment = align
        if space_after is not None and hasattr(x, "paragraph_format"): x.paragraph_format.space_after = Pt(space_after)
    for n in ["Normal", "Body Text", "First Paragraph", "Compact", "MDPI Back"]:
        font(n, 10, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    font("MDPI Type", 10, italic=True)
    font("MDPI Title", 18, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=12)
    font("MDPI Authors", 10, bold=True, space_after=8)
    font("MDPI Affil", 8.5, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
    font("MDPI Abstract", 9, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)
    font("MDPI Keywords", 9, space_after=6)
    for x in st:
        if x.name and x.name.lower().startswith("heading"):
            x.font.color.rgb = RGBColor(0, 0, 0); x.font.name = "Palatino Linotype"; x.font.size = Pt(10); x.font.bold = True
    font("Table Caption", 9); font("Image Caption", 9, italic=False)
    font("New Text", color=RGBColor(0x1F, 0x4E, 0xB4) if marked else RGBColor(0, 0, 0))
    font("Verify Text", bold=True, color=RGBColor(0xC0, 0x00, 0x00))
    for sec in d.sections:
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.0); sec.top_margin = sec.bottom_margin = Cm(2.0)
    # tables: small font + top/bottom rules
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    for t in d.tables:
        tblPr = t._tbl.tblPr
        b = OxmlElement("w:tblBorders")
        for edge in ("top", "bottom", "insideH"):
            e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "000000"); b.append(e)
        tblPr.append(b)
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs: r.font.size = Pt(8)
    d.save(dst)

style_doc("_raw.docx", "Grover_noise_manuscript_REVISED_marked.docx", True)
style_doc("_raw.docx", "Grover_noise_manuscript_REVISED_clean.docx", False)
print("done")
